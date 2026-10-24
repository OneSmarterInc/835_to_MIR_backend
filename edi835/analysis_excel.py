"""Plain, bounded-memory Excel export of all Analysis comparison categories."""
from decimal import Decimal
from math import ceil
from pathlib import Path
from tempfile import SpooledTemporaryFile
from typing import Any, BinaryIO

from django.db import connection
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU

from accounts.models import Client
from .analysis_queries import add_denial_codes, comparison_sql, dictionaries, json_values, summary_row
from .models import AnalysisUpload


HEADERS = (
    "Claim ID", "Internal claim number", "837 source", "Billed ($)",
    "Paid in supplied 835s ($)", "Unconfirmed / difference ($)",
    "Remit records", "835 sources", "Notes",
)
SHEETS = ("Mismatch amount", "Matched amount", "Missing in 835")
COLUMN_WIDTHS = (26, 24, 42, 19, 25, 28, 16, 42, 52)
DENIAL_HEADERS = ("Denial code", "Denial reason")
DENIAL_WIDTHS = (18, 80)
CURRENCY = '"$"#,##0.00;-"$"#,##0.00'


def comparison_workbook(client: Client, options: dict[str, Any]) -> BinaryIO:
    """Export all categories from one query snapshot, ignoring UI pagination.

    Options must come from query_options. The caller owns the returned file and
    must close it after the response. Temporary output spills to disk above 8 MiB.
    """
    workbook = Workbook(write_only=True)
    workbook.properties.creator = "Filex OS"
    workbook.properties.lastModifiedBy = "Filex OS"
    workbook.properties.title = "FileX Claims Reconciliation"
    workbook.properties.subject = "Analysis allowed comparisons"
    workbook.properties.description = "FileX Analysis export: 837 billed charges, cumulative payments and allowed comparison."
    workbook.properties.keywords = "FileX, Filex OS, Analysis, 835, 837, reconciliation"
    workbook.properties.category = "FileX Analysis"
    workbook.properties.created = workbook.properties.modified = timezone.now().replace(tzinfo=None)
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    header_font = Font(bold=True)
    wrapped = Alignment(vertical="top", wrap_text=True)
    edge = Side(style="thin", color="B8C2CC")
    grid = Border(left=edge, right=edge, top=edge, bottom=edge)
    sheets = []
    counts = [3, 3, 3]

    def cell(sheet, value, *, currency=False):
        result = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            # IDs are literal text, including leading zeroes and formula prefixes.
            result.data_type = "s"
        result.alignment = wrapped
        if currency:
            result.number_format = CURRENCY
        return result

    try:
        for title in SHEETS:
            sheet = workbook.create_sheet(title)
            sheets.append(sheet)
            sheet.freeze_panes = "C4"
            sheet.sheet_view.showGridLines = False
            sheet.sheet_format.defaultRowHeight = 44
            sheet.row_dimensions[1].height = 34
            sheet.row_dimensions[2].height = 24
            sheet.row_dimensions[3].height = 46
            for index, width in enumerate(COLUMN_WIDTHS + (DENIAL_WIDTHS if title == SHEETS[0] else ()), 1):
                sheet.column_dimensions[get_column_letter(index)].width = width
            # Transparent portal mark and a reserved three-character logo slot.
            # Literal spacing also renders correctly in viewers that ignore indent.
            logo = Image(Path(__file__).parent / "assets" / "filex-mark.png")
            logo.anchor = OneCellAnchor(
                _from=AnchorMarker(col=0, row=0, colOff=pixels_to_EMU(2), rowOff=pixels_to_EMU(2)),
                ext=XDRPositiveSize2D(pixels_to_EMU(40), pixels_to_EMU(40)),
            )
            sheet.add_image(logo)
            wordmark = cell(sheet, "   FileX")
            wordmark.font = Font(name="Consolas", size=18, bold=True, color="172033")
            wordmark.alignment = Alignment(horizontal="left", vertical="center")
            report_title = cell(sheet, "Claims reconciliation")
            report_title.alignment = Alignment(vertical="center")
            sheet.append([wordmark, None, report_title])
            section = cell(sheet, title)
            section.font = Font(size=12, bold=True)
            sheet.append([section])
            headers = [cell(sheet, heading) for heading in HEADERS + (DENIAL_HEADERS if title == SHEETS[0] else ())]
            for header in headers:
                header.font = header_font
                header.border = grid
            sheet.append(headers)

        has835 = AnalysisUpload.objects.filter(client=client, kind="835").exists()
        sql, params = comparison_sql(client, options["search"])
        # Identifiers are validated by query_options; no raw request SQL.
        order = f" ORDER BY {options['sort']} {options['direction'].upper()} NULLS LAST, claim ASC"
        with connection.cursor() as cursor:
            cursor.execute(sql + "SELECT * FROM compared" + order, params)
            while batch := cursor.fetchmany(500):
                records = dictionaries(cursor, batch)
                rows = [summary_row(record, has835) for record in records]
                add_denial_codes(client, [row for row in rows if row["category"] == "both" and row["allowed_status"] != "Match"])
                for record, row in zip(records, rows):
                    matched = row["allowed_status"] == "Match"
                    index = 2 if record["n835"] == 0 else 1 if matched else 0
                    counts[index] += 1
                    if counts[index] > 1048576:
                        raise ValueError("Too many claims for one Excel sheet. Narrow the search and retry.")
                    sheet = sheets[index]
                    reference = "\n".join(row["internal_numbers"]) or "—"
                    billed = Decimal(row["billed_837"]) if row["billed_837"] is not None else None
                    payment = Decimal(row["payment"]) if row["payment"] is not None else None
                    # Template difference is billed minus paid, not the allowed comparison delta.
                    difference = (billed if record["n835"] == 0 else billed - payment
                                  if billed is not None and payment is not None else None)
                    if record["n835"] == 0:
                        note = "No matching CLP01 in supplied 835 files; payment unconfirmed."
                    elif row["allowed_status"] == "Match":
                        note = ("Payments equal original billed amount." if row["allowed_basis"] == "payment"
                                else "Payments plus patient responsibility equal original billed amount.")
                    else:
                        note = f"Allowed comparison: {row['allowed_status']}."
                        if row["allowed_difference"] is not None:
                            note += f" Difference ${Decimal(row['allowed_difference']):,.2f}."
                    raw_values = [row["claim"], reference, "\n".join(json_values(record["sources_837"])),
                                  billed, payment, difference, record["n835"],
                                  "\n".join(json_values(record["sources_835"])) or "—", note]
                    if index == 0:
                        raw_values.extend([
                            "\n".join(row["denial_codes"]) or "—",
                            "\n".join(item["reason"] or "—" for item in row["denial_reasons"]) or "—",
                        ])
                    values = [cell(sheet, value if value is not None else "—", currency=isinstance(value, Decimal))
                              for value in raw_values]
                    lines = max(sum(max(1, ceil(len(line) / (width - 2))) for line in str(value or "").split("\n"))
                                for value, width in zip(raw_values, COLUMN_WIDTHS + (DENIAL_WIDTHS if index == 0 else ())))
                    sheet.row_dimensions[counts[index]].height = min(409, max(44, 15 * lines + 8))
                    for value_cell in values:
                        value_cell.border = grid
                    sheet.append(values)
        for sheet, count in zip(sheets, counts):
            column_count = len(HEADERS) + (len(DENIAL_HEADERS) if sheet.title == SHEETS[0] else 0)
            sheet.auto_filter.ref = f"A3:{get_column_letter(column_count)}{count}"
        workbook.save(output)
        output.seek(0)
        return output
    except BaseException:
        output.close()
        raise
    finally:
        workbook.close()
