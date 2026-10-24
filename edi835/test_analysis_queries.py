import csv
import json
import zipfile
from decimal import Decimal
from io import BytesIO, StringIO
from time import perf_counter
from unittest import skipUnless

from django.db import connection
from django.test import RequestFactory, TestCase
from django.test.utils import CaptureQueriesContext

from accounts.models import Client, User
from .analysis_queries import claim_detail, comparison_csv, comparison_page, query_options
from .client_portal_analysis import analysis_upload
from .models import AnalysisClaim, AnalysisClaimRejection, AnalysisUpload, Analysis835Claim, ClaimAdjustmentReason


class AnalysisQueryTests(TestCase):
    def setUp(self):
        self.client_record = Client.objects.create(name="Query", client_code="QUERY", email="query@example.com")
        self.user = User.objects.create_user(email="queries@example.com", name="Viewer", mobile="15555550235",
                                            client=self.client_record, client_role="CLIENT_USER", password="TestPassword!234")
        self.counter = 0

    def upload(self, kind, claims, filename=None, tenant=None, paid_on=None):
        self.counter += 1
        upload = AnalysisUpload.objects.create(client=tenant or self.client_record, kind=kind,
            filename=filename or f"file-{self.counter}.{kind}", content_sha256=str(self.counter), raw_content=b"private",
            accepted_count=len(claims), processing_version=2)
        for index, values in enumerate(claims, 1):
            number, billed, payment, responsibility, internal, claim_status = (*values, *([None] * (6 - len(values))))
            projection = AnalysisClaim.objects.create(upload=upload, sequence=index, claim_number=number, billed=Decimal(billed),
                payment=payment, patient_responsibility=responsibility, internal_number=internal or "REF", payment_date=paid_on)
            if kind == "835":
                Analysis835Claim.objects.create(edi_file=upload, analysis_claim=projection, claim_sequence=index,
                    highmark_claim_number=number, internal_claim_number=internal or "REF", claim_status=claim_status or "1",
                    total_charge_amount=billed, paid_amount=payment or "0", patient_responsibility=responsibility or "0")
        return upload

    def page(self, **options):
        return comparison_page(self.client_record, query_options(options))

    def get(self, **params):
        request = RequestFactory().get("/edi835/api/client-portal/analysis/", params)
        request.user = self.user
        return analysis_upload(request)

    def test_denial_codes_include_claim_and_service_cas_without_tenant_leakage(self):
        from importlib import import_module
        from django.apps import apps
        seed = import_module("edi835.migrations.0070_claim_adjustment_reasons").seed_reasons
        from types import SimpleNamespace
        seed(apps, SimpleNamespace(connection=connection))
        seed(apps, SimpleNamespace(connection=connection))
        self.assertEqual(ClaimAdjustmentReason.objects.count(), 407)
        self.upload("837", [("CAS-CLAIM", "100"), ("MISSING", "100")])
        source = self.upload("835", [(" cas-claim ", "100", "0", "0")])
        source.normalized_835_claims.update(segment_data={
            "element_separator": "|", "segments": [
                "CLP|CAS-CLAIM|4|100|0", "CAS|CO|45|10||97|20|1",
                "SVC|HC:99213|70|0", "CAS|CO|45|10", "CAS|PR|1|60",
            ]})
        other = Client.objects.create(name="Other", client_code="CAS-OTHER", email="cas-other@example.com")
        foreign = self.upload("835", [("CAS-CLAIM", "100", "0", "0")], tenant=other)
        foreign.normalized_835_claims.update(segment_data={"segments": ["CAS*OA*23*100"]})
        row = self.page()["rows"][0]
        self.assertEqual(row["denial_codes"], ["CO-45", "CO-97", "PR-1"])
        self.assertEqual(row["denial_reasons"][-1], {"code": "PR-1", "reason": "Deductible Amount"})
        self.assertEqual(row["denial_reasons"][0]["reason"], ClaimAdjustmentReason.objects.get(pk="45").description)
        self.assertEqual(claim_detail(self.client_record, "CAS-CLAIM")["denial_codes"], row["denial_codes"])
        self.assertEqual(self.page(category="837-only")["rows"][0]["denial_codes"], [])
        source.normalized_835_claims.update(segment_data={})
        self.assertEqual(self.page()["rows"][0]["denial_codes"], [])
        source.normalized_835_claims.update(segment_data={"segments": ["CAS*CO*B11*10", "CAS*OA*99999*90"]})
        reasons = self.page()["rows"][0]["denial_reasons"]
        self.assertEqual(reasons, [
            {"code": "CO-B11", "reason": ClaimAdjustmentReason.objects.get(pk="B11").description},
            {"code": "OA-99999", "reason": None},
        ])

    def test_signed_comparisons_unknowns_and_complete_group_pagination(self):
        self.upload("837", [("MATCH", "657"), ("BILLED", "100"), ("ALLOWED", "657"), ("SIGNED", "100"),
                            ("UNKNOWN", "100"), ("MISSING", "1"), ("CENTS", ".30")])
        self.upload("835", [("MATCH", "657", "0", "657"), ("BILLED", "80", "70", "30"),
                            ("ALLOWED", "657", "600", "0"), ("SIGNED", "-100", "-100", "0"),
                            ("UNKNOWN", "100", "10", None), ("CENTS", ".30", ".10", ".20")])
        result = self.page()
        self.assertEqual(result["counts"], {"both": 3, "matched": 3, "837-only": 1})
        rows = {row["claim"]: row for row in result["rows"]}
        self.assertEqual(rows["SIGNED"]["difference"], "200.00")
        self.assertNotIn("BILLED", rows)
        self.assertEqual(rows["ALLOWED"]["billed_status"], "Match")
        self.assertEqual(rows["UNKNOWN"]["allowed_status"], "Amount unavailable")
        self.assertEqual(rows["UNKNOWN"]["occurrences"], [])
        self.assertNotIn("filename", rows["ALLOWED"])
        self.assertEqual(self.page(category="837-only")["rows"][0]["claim"], "MISSING")
        matched = self.page(category="matched")
        self.assertEqual({row["claim"] for row in matched["rows"]}, {"MATCH", "CENTS", "BILLED"})
        self.assertTrue(all(row["allowed_status"] == "Match" for row in matched["rows"]))

    def test_matched_category_pagination_search_export_and_ambiguity(self):
        self.upload("837", [(f"M{i:02}", str(i)) for i in range(23)] + [("NEG", "-100"), ("AMB", "50")])
        self.upload("835", [(f"M{i:02}", str(i), str(i), "0") for i in range(23)] +
                    [("NEG", "-100", "-80", "-20"), ("AMB", "50", "50", "0")])
        self.upload("835", [("AMB", "50", "50", "0")], "different-source.835")
        page = self.page(category="matched", page="2", sort="claim", direction="desc")
        self.assertEqual((page["total"], page["pages"], len(page["rows"])), (24, 3, 10))
        self.assertNotIn("AMB", [row["claim"] for row in page["rows"]])
        self.assertEqual(self.page(category="matched", search="NEG")["rows"][0]["difference"], "0.00")
        self.assertEqual(self.page(category="matched", search="absent")["total"], 0)
        response = self.get(view="export", category="matched", search="M0", page="2")
        self.assertEqual(response.status_code, 200)
        self.assertIn("matched-amount.csv", response["Content-Disposition"])
        rows = list(csv.reader(StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
        self.assertEqual(len(rows), 11)
        self.assertTrue(all(row[12] == row[14] == "Match" for row in rows[1:]))
        other = Client.objects.create(name="Other", client_code="MATCH-OTHER", email="matched-other@example.com")
        self.upload("837", [("FOREIGN", "1")], tenant=other)
        self.upload("835", [("FOREIGN", "1", "1", "0")], tenant=other)
        self.assertEqual(self.page(category="matched")["total"], 24)

    def test_distinct_filename_financial_records_and_display_values(self):
        self.upload("837", [(" claim ", "100", None, None, " r "), ("CLAIM", "100", None, None, "R")], "same.837")
        self.upload("837", [("CLAIM", "100", None, None, "R")], "same.837")
        self.upload("835", [("CLAIM", "80", "50", "20", "R")], "same.835")
        row = self.page()["rows"][0]
        self.assertFalse(row["ambiguous"])
        self.assertEqual(row["difference"], "20.00")
        self.assertEqual(len(claim_detail(self.client_record, "claim")["occurrences"]), 2)
        self.upload("835", [("CLAIM", "90", "50", "20", "OTHER")], "other.835")
        row = self.page(category="matched")["rows"][0]
        self.assertFalse(row["ambiguous"])
        self.assertEqual(row["reference_count"], 2)
        self.assertEqual(row["display_amounts"]["billed_835"]["count"], 1)
        self.assertEqual(row["display_amounts"]["payment"], {"count": 1, "value": "100.00"})
        self.assertEqual(self.page(category="matched")["total"], 1)
        self.assertEqual(len(claim_detail(self.client_record, row["key"])["occurrences"]), 3)

    def test_search_keeps_entire_group_and_treats_wildcards_literally(self):
        self.upload("837", [("X", "100", None, None, "837-REF")], "percent%_file.837")
        self.upload("835", [("X", "80", "50", "0", "835-REF")], "response.835")
        for search in ("837-ref", "835-ref", "percent%_", "response", "X"):
            row = self.page(search=search)["rows"][0]
            self.assertEqual(row["category"], "both")
            self.assertEqual(row["difference"], "20.00")
        self.assertEqual(self.page(search="notfound")["total"], 0)
        self.assertEqual(self.page(search="' OR 1=1 --")["total"], 0)
        self.assertEqual(self.page(search="%%%%")["total"], 0)

    def test_835_totals_drive_both_comparisons_counts_detail_and_csv(self):
        self.upload("837", [("TOTAL", "624")])
        self.upload("835", [("TOTAL", "400", "350", "50", "REF")], "first.835", paid_on="2026-09-01")
        self.upload("835", [("TOTAL", "624", "200", "24", "REF")], "second.835", paid_on="2026-09-02")
        result = self.page(category="matched", search="second.835")
        self.assertEqual(result["counts"], {"both": 0, "matched": 1, "837-only": 0})
        row = result["rows"][0]
        self.assertEqual([row[key] for key in ("billed_835", "payment", "patient_responsibility", "allowed")],
                         ["624.00", "550.00", "74.00", "624.00"])
        self.assertEqual((row["billed_status"], row["allowed_status"]), ("Match", "Match"))
        self.assertFalse(row["ambiguous"])
        self.assertEqual(row["reference_count"], 1)
        self.assertEqual(claim_detail(self.client_record, "TOTAL")["billed_835"], "624.00")
        self.assertEqual(len(claim_detail(self.client_record, "TOTAL")["occurrences"]), 3)
        exported = list(csv.reader(StringIO("".join(comparison_csv(
            self.client_record, query_options({"category": "matched"}))))))
        self.assertEqual(exported[1][7:11], ["624.00", "550.00", "74.00", "624.00"])

    def excel(self, **params):
        from openpyxl import load_workbook
        response = self.get(view="excel", **params)
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response["Cache-Control"])
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn('filename="Claims-Reconciliation.xlsx"', response["Content-Disposition"])
        try:
            content = b"".join(response.streaming_content)
            with zipfile.ZipFile(BytesIO(content)) as archive:
                self.assertEqual(len([name for name in archive.namelist() if name.startswith("xl/media/")]), 3)
                for number in range(1, 4):
                    self.assertIn(b"<drawing", archive.read(f"xl/worksheets/sheet{number}.xml"))
                    from PIL import Image
                    logo = Image.open(BytesIO(archive.read(f"xl/media/image{number}.png")))
                    self.assertEqual(logo.mode, "RGBA")
                    self.assertEqual(logo.getpixel((0, 0))[3], 0)
            return load_workbook(BytesIO(content))
        finally:
            response.close()

    def test_excel_contains_all_categories_numeric_totals_and_table_headings(self):
        from .analysis_excel import HEADERS, SHEETS
        self.assertEqual(HEADERS, ("Claim ID", "Internal claim number", "837 source", "Billed ($)",
                                   "Paid in supplied 835s ($)", "Unconfirmed / difference ($)",
                                   "Remit records", "835 sources", "Notes"))
        self.upload("837", [("00123", "100"), ("MISMATCH", "100"), ("MISSING", "39")])
        self.upload("835", [("00123", "120", "100", "20"), ("MISMATCH", "80", "50", "20")], paid_on="2026-09-01")
        self.upload("835", [("00123", "100", "-10", "-10")], paid_on="2026-09-02")
        # UI selection and pagination must not suppress other export sheets.
        with CaptureQueriesContext(connection) as queries:
            workbook = self.excel(category="837-only", page="999", page_size="10")
        self.assertEqual(len(queries), 2)
        self.assertEqual(workbook.sheetnames, list(SHEETS))
        self.assertEqual(workbook.properties.creator, "Filex OS")
        self.assertEqual(workbook.properties.lastModifiedBy, "Filex OS")
        self.assertEqual(workbook.properties.title, "FileX Claims Reconciliation")
        for sheet in workbook:
            self.assertEqual(sheet["A1"].value.strip(), "FileX")
            self.assertEqual(tuple(cell.value for cell in sheet[3]), HEADERS)
            self.assertEqual(sheet.max_row, 4)
            self.assertEqual(sheet.max_column, 9)
            self.assertEqual(sheet.freeze_panes, "C4")
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertTrue(all(cell.border.left.style == "thin" and cell.border.bottom.style == "thin"
                                for row in sheet.iter_rows(min_row=3) for cell in row))
            self.assertIsNone(sheet["A1"].border.left.style)
            self.assertEqual(sheet.auto_filter.ref, "A3:I4")
            self.assertTrue(all(cell.fill.patternType is None for row in sheet.iter_rows(min_row=3) for cell in row))
            self.assertTrue(all(cell.data_type != "f" for row in sheet for cell in row))
        matched = workbook["Matched amount"]
        self.assertEqual(matched["A4"].value, "00123")
        self.assertEqual(matched["A4"].data_type, "s")
        self.assertEqual([matched.cell(4, col).value for col in range(4, 8)], [100, 90, 10, 2])
        self.assertTrue(all(matched.cell(4, col).data_type == "n" for col in range(4, 8)))
        self.assertEqual(matched["I4"].value, "Payments plus patient responsibility equal original billed amount.")
        self.assertEqual(workbook["Mismatch amount"]["I4"].value, "Allowed comparison: Mismatch. Difference $30.00.")
        self.assertEqual(workbook["Missing in 835"]["E4"].value, "—")
        self.assertEqual(workbook["Missing in 835"]["F4"].value, 39)
        self.assertEqual(workbook["Missing in 835"]["I4"].value, "No matching CLP01 in supplied 835 files; payment unconfirmed.")
        workbook.close()

    def test_excel_exports_all_pages_search_tenant_boundary_and_literal_ids(self):
        self.upload("837", [(f"C{i:03}", str(i)) for i in range(12)] + [("=FORMULA", "-20")])
        self.upload("835", [("=FORMULA", "-10", "-10", "0", "=INTERNAL")])
        other = Client.objects.create(name="Other Excel", client_code="OTHERXL", email="otherxl@example.com")
        self.upload("837", [("SECRET", "50")], tenant=other)
        workbook = self.excel(category="both", sort="billed_837", direction="desc", page="2")
        self.assertEqual(workbook["Missing in 835"].max_row, 15)
        self.assertEqual(workbook["Missing in 835"]["A4"].value, "C011")
        sheet = workbook["Mismatch amount"]
        self.assertEqual((sheet["A4"].value, sheet["A4"].data_type), ("=FORMULA", "s"))
        self.assertEqual((sheet["B4"].value, sheet["B4"].data_type), ("=INTERNAL", "s"))
        self.assertEqual(sheet["D4"].value, -20)
        self.assertNotIn("SECRET", [cell.value for tab in workbook for row in tab for cell in row])
        workbook.close()
        filtered = self.excel(search="C011")
        self.assertEqual([sheet.max_row for sheet in filtered], [3, 3, 4])
        filtered.close()
        empty = self.excel(search="ABSENT")
        self.assertEqual([sheet.max_row for sheet in empty], [3, 3, 3])
        empty.close()

    def test_sum_preserves_reversals_unknowns_and_same_file_deduplication(self):
        self.upload("837", [("SIGNED", "100"), ("UNKNOWN", "100"), ("EQUAL", "100")])
        self.upload("835", [("SIGNED", "120", "100", "20"), ("UNKNOWN", "50", "30", "20"),
                            ("EQUAL", "100", "40", "10")], "first.835", paid_on="2026-09-01")
        self.upload("835", [("SIGNED", "100", "-10", "-10"), ("UNKNOWN", "100", "50", None),
                            ("EQUAL", "100", "40", "10")], "second.835", paid_on="2026-09-02")
        # A repeated record in the same named file remains a single contribution.
        self.upload("835", [("SIGNED", "100", "-10", "-10")], "second.835", paid_on="2026-09-02")
        matched = {row["claim"]: row for row in self.page(category="matched")["rows"]}
        self.assertEqual(set(matched), {"SIGNED", "EQUAL"})
        self.assertEqual((matched["SIGNED"]["payment"], matched["SIGNED"]["patient_responsibility"]),
                         ("90.00", "10.00"))
        self.assertEqual(matched["EQUAL"]["payment"], "80.00")
        unknown = self.page()["rows"][0]
        self.assertEqual((unknown["billed_status"], unknown["allowed_status"]), ("Match", "Amount unavailable"))
        self.assertEqual(unknown["payment"], "80.00")
        self.assertIsNone(unknown["patient_responsibility"])
        self.assertIsNone(unknown["allowed"])

    def test_latest_upload_overrides_payment_date_and_payments_can_match_alone(self):
        self.upload("837", [("LATEST", "624")])
        self.upload("835", [("LATEST", "117", "117", "0", " ref ")], "z.835", paid_on="2026-09-16")
        latest = self.upload("835", [("LATEST", "624", "507", "20", "REF")], "a.835", paid_on="2026-09-08")
        row = self.page(category="matched")["rows"][0]
        self.assertEqual((row["billed_835"], row["payment"], row["allowed"]), ("624.00", "624.00", "644.00"))
        self.assertEqual((row["billed_status"], row["allowed_status"], row["allowed_difference"]), ("Match", "Match", "0.00"))
        self.assertEqual(row["allowed_basis"], "payment")
        self.assertEqual(row["latest_payment_date"], "2026-09-08")
        self.assertEqual((row["reference_count"], row["internal_numbers"]), (1, ["REF"]))
        self.assertEqual(claim_detail(self.client_record, "LATEST")["billed_835"], "624.00")
        exported = self.excel()
        self.assertEqual(exported["Matched amount"]["D4"].value, 624)
        self.assertEqual(exported["Matched amount"]["B4"].value, "REF")
        self.assertEqual(exported["Matched amount"]["I4"].value, "Payments equal original billed amount.")
        exported.close()
        # Changing BPR16 cannot change the selected uploaded file.
        latest.claims.update(payment_date=None)
        self.assertEqual(self.page(category="matched")["rows"][0]["billed_835"], "624.00")

    def test_latest_clp03_is_neither_sum_nor_largest_amount(self):
        self.upload("837", [("X", "4")])
        self.upload("835", [("X", "1", "1", "0", "REF")])
        self.upload("835", [("X", "4", "3", "0", "REF")])
        row = self.page(category="matched")["rows"][0]
        self.assertEqual((row["billed_835"], row["payment"]), ("4.00", "4.00"))
        self.assertEqual(row["internal_numbers"], ["REF"])
        exported = self.excel()
        self.assertEqual(exported["Matched amount"]["B4"].value, "REF")
        exported.close()
        self.upload("835", [("X", "2", "0", "0", "REF")])
        self.assertEqual(self.page(category="matched")["rows"][0]["billed_835"], "2.00")

    def test_clp_sequence_selects_actual_latest_value_not_sum_or_maximum(self):
        self.upload("837", [("X", "60")])
        self.upload("835", [("X", "30", "30", "0"), ("X", "60", "30", "0")])
        row = self.page(category="matched")["rows"][0]
        self.assertEqual((row["billed_835"], row["payment"]), ("60.00", "60.00"))
        self.upload("835", [("X", "90", "0", "0"), ("X", "40", "0", "0")])
        self.assertEqual(self.page(category="matched")["rows"][0]["billed_835"], "40.00")

    def test_reversal_and_correction_use_status_and_sequence(self):
        self.upload("837", [("X", "60")])
        self.upload("835", [("X", "30", "30", "0", "REF", "1")])
        self.upload("835", [("X", "-30", "-30", "0", "REF", "22")])
        row = self.page()["rows"][0]
        self.assertEqual(row["latest_claim_status"], "22")
        self.assertEqual((row["billed_835"], row["payment"]), ("-30.00", "0.00"))
        self.assertEqual(row["billed_status"], "Reversed — review")
        self.assertEqual(row["allowed_status"], "Reversed — review")
        self.upload("835", [("X", "60", "60", "0", "REF", "1")])
        row = self.page(category="matched")["rows"][0]
        self.assertEqual((row["billed_835"], row["payment"]), ("60.00", "60.00"))
        detail = claim_detail(self.client_record, row["key"])
        self.assertEqual(detail["selected_record_id"], row["selected_record_id"])
        self.assertIn("22", [item["claim_status"] for item in detail["occurrences"]])

    def test_internal_claims_share_one_row_and_allowed_match_controls_export(self):
        self.upload("837", [("X", "624")], "submitted.837")
        self.upload("835", [("X", "624", "507", "0", "FIRST0000")], "first.835")
        self.upload("835", [("X", "117", "117", "0", "SECOND0000")], "second.835")
        result = self.page(category="matched")
        self.assertEqual(result["total"], 1)
        row = result["rows"][0]
        self.assertEqual(row["internal_numbers"], ["FIRST0000", "SECOND0000"])
        self.assertEqual((row["payment"], row["billed_status"], row["allowed_status"]), ("624.00", "Mismatch", "Match"))
        self.assertEqual(self.page()["total"], 0)
        response = self.get(view="claim", claim=row["key"])
        self.assertEqual(response.status_code, 200)
        items = json.loads(response.content)["row"]["occurrences"]
        self.assertEqual({item["internal_number"] for item in items if item["source_type"] == "835"}, {"FIRST0000", "SECOND0000"})
        self.assertEqual(len(claim_detail(self.client_record, '["X","id:FIRST0000"]')["occurrences"]), 3)
        self.assertEqual(self.page(search="SECOND0000", category="matched")["total"], 1)
        exported = self.excel()
        sheet = exported["Matched amount"]
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual([sheet.cell(4, col).value for col in range(1, 10)],
                         ["X", "FIRST0000\nSECOND0000", "submitted.837", 624, 624, 0, 2,
                          "first.835\nsecond.835", "Payments equal original billed amount."])
        self.assertEqual(exported["Mismatch amount"].max_row, 3)
        exported.close()

    def test_missing_internal_ids_and_different_patient_claims_do_not_merge(self):
        self.upload("837", [("X", "60"), ("Y", "60")])
        unknown = self.upload("835", [("X", "60", "30", "0"), ("X", "60", "30", "0")])
        unknown.claims.update(internal_number="")
        self.upload("835", [("Y", "60", "60", "0", "REF")])
        self.upload("835", [("X", "60", "60", "0", "REF")])
        rows = self.page()["rows"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["claim"], "X")
        for row in rows:
            self.assertEqual(row["payment"], "120.00")
            self.assertEqual(len([item for item in claim_detail(self.client_record, row["key"])["occurrences"]
                                  if item["source_type"] == "835"]), 3)
        self.assertEqual(self.page(category="matched")["rows"][0]["claim"], "Y")

    def test_final_reversal_in_any_internal_claim_remains_review(self):
        self.upload("837", [("X", "100")])
        self.upload("835", [("X", "50", "50", "0", "A")])
        self.upload("835", [("X", "-50", "-50", "0", "A", "22")])
        self.upload("835", [("X", "100", "100", "0", "B")])
        row = self.page()["rows"][0]
        self.assertEqual(row["payment"], "100.00")
        self.assertEqual(row["allowed_status"], "Reversed — review")
        self.assertEqual(self.page(category="matched")["total"], 0)
        self.upload("835", [("X", "0", "0", "0", "A", "1")])
        self.assertEqual(self.page(category="matched")["total"], 1)

    def test_internal_claim_is_from_835_only_in_list_detail_and_export(self):
        self.upload("837", [("X", "100", None, None, "SUBMISSION-REF")])
        self.upload("835", [("X", "80", "50", "0", "PAYER-REF")])
        row = self.page()["rows"][0]
        self.assertEqual((row["internal_number"], row["reference_count"]), ("PAYER-REF", 1))
        detail = claim_detail(self.client_record, "X")
        self.assertEqual([item["internal_number"] for item in detail["occurrences"] if item["source_type"] == "837"], [""])
        records = list(csv.reader(StringIO("".join(comparison_csv(self.client_record, query_options({"category": "both"}))))))
        self.assertEqual(records[1][5], "PAYER-REF")
        self.upload("837", [("MISSING", "100", None, None, "SUBMISSION-ONLY")])
        missing = self.page(category="837-only")["rows"][0]
        self.assertEqual((missing["internal_number"], missing["reference_count"]), ("", 0))
        self.upload("835", [("X", "80", "50", "0", "OTHER-PAYER-REF")])
        self.assertEqual(self.page(category="matched")["total"], 1)
        self.assertEqual(self.page(category="matched")["rows"][0]["reference_count"], 2)

    def test_constant_queries_bounded_payload_paging_sort_and_full_csv(self):
        self.upload("837", [(f"C{i:03}", str(i)) for i in range(51)])
        self.upload("835", [(f"C{i:03}", "99", "90", "0") for i in range(51)])
        with CaptureQueriesContext(connection) as queries:
            result = self.page(page="2", sort="billed_837", direction="desc")
        self.assertEqual(len(queries), 3)
        self.assertEqual(len(result["rows"]), 10)
        self.assertEqual(result["total"], 51)
        self.assertEqual(result["rows"][0]["claim"], "C040")
        self.assertTrue(all(not row["occurrences"] for row in result["rows"]))
        self.assertFalse(any('"raw_content"' in query["sql"] for query in queries))
        self.assertEqual(self.page(page="999")["page"], 6)
        with CaptureQueriesContext(connection) as export_queries:
            text = "".join(comparison_csv(self.client_record, query_options({"category": "both", "page": "2"})))
        self.assertEqual(len(export_queries), 3, "Export batches sources without per-claim queries")
        records = list(csv.reader(StringIO(text.lstrip("\ufeff"))))
        self.assertEqual(len(records), 52)
        self.assertEqual(records[1][6:11], ["0.00", "99.00", "90.00", "0.00", "90.00"])
        text = "".join(comparison_csv(self.client_record, query_options({"category": "both", "search": "C050"})))
        self.assertEqual(len(list(csv.reader(StringIO(text)))), 2)

    def test_private_reads_lazy_reports_and_input_validation(self):
        other = Client.objects.create(name="Other", client_code="OTHERQ", email="otherq@example.com")
        foreign = self.upload("837", [("SECRET", "100")], tenant=other)
        local = self.upload("837", [("LOCAL", "100")])
        self.upload("835", [("LOCAL", "100", "100", "0")], tenant=other)
        AnalysisClaimRejection.objects.create(upload=local, sequence=3, claim_number="BAD", reason="Invalid amount")
        local.rejected_count = 1
        local.save(update_fields=["rejected_count"])
        response = self.get(view="comparison")
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response["Cache-Control"])
        self.assertNotIn("SECRET", response.content.decode())
        self.assertNotIn("Invalid amount", response.content.decode())
        self.assertEqual(json.loads(response.content)["counts"], {"both": 0, "matched": 0, "837-only": 1})
        self.assertEqual(self.get(view="claim", claim="SECRET").status_code, 404)
        self.assertEqual(self.get(view="rejections", upload=str(foreign.pk)).status_code, 404)
        self.assertEqual(self.get(view="rejection-export", upload=str(foreign.pk)).status_code, 404)
        result = json.loads(self.get(view="rejections", upload=str(local.pk)).content)
        self.assertEqual(result["rejections"][0]["reason"], "Invalid amount")
        self.assertTrue(self.get(view="rejection-export", upload=str(local.pk)).streaming)
        self.assertTrue(self.get(view="export", category="837-only").streaming)
        for invalid in ({"sort": "claim; DROP TABLE"}, {"page": "0"}, {"page_size": "999"}, {"direction": "invalid"}, {"search": "x" * 101}):
            self.assertEqual(self.get(view="comparison", **invalid).status_code, 400)
        detail = json.loads(self.get(view="claim", claim="LOCAL").content)["row"]
        self.assertEqual(detail["occurrences"][0]["filename"], local.filename)

    def test_empty_dataset_and_missing_status(self):
        self.assertEqual(self.page()["rows"], [])
        self.upload("837", [("A", "100")])
        self.assertEqual(self.page()["rows"][0]["billed_status"], "Awaiting 835 upload")
        self.upload("835", [])
        self.assertEqual(self.page(category="837-only")["rows"][0]["billed_status"], "Not in uploaded 835")

    @skipUnless(connection.vendor == "postgresql", "Exact maximum NUMERIC precision requires PostgreSQL")
    def test_postgres_maximum_precision_and_sign(self):
        self.upload("837", [("CENT", "9999999999999999.99"), ("NEGATIVE", "9999999999999999.99")])
        self.upload("835", [("CENT", "9999999999999999.98", "9999999999999999.98", "0"),
                            ("NEGATIVE", "-9999999999999999.99", "-9999999999999999.99", "-9999999999999999.99")])
        rows = {row["claim"]: row for row in self.page()["rows"]}
        self.assertEqual(rows["CENT"]["difference"], "0.01")
        self.assertEqual(rows["NEGATIVE"]["allowed_difference"], "29999999999999999.97")

    def test_large_history_request_and_payload_reduction(self):
        from .client_portal_analysis import upload_payload
        uploads = [AnalysisUpload(client=self.client_record, kind=kind, filename=f"{index}.{kind}",
                   content_sha256=str(index), raw_content=b"private", accepted_count=200, processing_version=2)
                   for kind in ("837", "835") for index in range(50)]
        AnalysisUpload.objects.bulk_create(uploads)
        claims = []
        for upload in uploads:
            index = int(upload.filename.split(".")[0])
            for sequence in range(200):
                claims.append(AnalysisClaim(upload=upload, sequence=sequence + 1, claim_number=f"C{index * 200 + sequence:05}",
                    internal_number="REF", billed=Decimal("100" if upload.kind == "837" else "80"),
                    payment=Decimal("70") if upload.kind == "835" else None,
                    patient_responsibility=Decimal("10") if upload.kind == "835" else None))
        AnalysisClaim.objects.bulk_create(claims, batch_size=1000)
        start = perf_counter()
        response = self.get(view="comparison")
        elapsed = perf_counter() - start
        result = json.loads(response.content)
        self.assertEqual(result["total"], 10000)
        self.assertEqual(len(result["rows"]), 10)
        old_bytes = sum(len(json.dumps(upload_payload(upload)).encode()) for upload in uploads)
        new_bytes = len(response.content)
        self.assertLess(new_bytes, old_bytes // 100)
        print(f"\nAnalysis benchmark ({connection.vendor}, synthetic): 100 files / 20,000 occurrences; "
              f"100 history requests -> 1 page request; {old_bytes:,} -> {new_bytes:,} payload bytes; "
              f"new request {elapsed:.3f}s.")

    def test_csv_injection_missing_status_and_limited_rejection_preview(self):
        upload = self.upload("837", [("=FORMULA", "10")], "+filename.837")
        exported = "".join(comparison_csv(self.client_record, query_options({"category": "837-only"})))
        self.assertIn("'=FORMULA", exported)
        self.assertIn("'+filename.837", exported)
        self.assertIn("Awaiting 835 upload", exported)
        AnalysisClaimRejection.objects.bulk_create([AnalysisClaimRejection(upload=upload, sequence=i,
            reason=f"Invalid amount {i}") for i in range(25)])
        upload.rejected_count = 25
        upload.save(update_fields=["rejected_count"])
        preview = json.loads(self.get(view="rejections", upload=str(upload.pk)).content)
        self.assertEqual(len(preview["rejections"]), 20)
        self.assertEqual(preview["total"], 25)
        report = self.get(view="rejection-export", upload=str(upload.pk))
        text = b"".join(report.streaming_content).decode("utf-8-sig")
        self.assertEqual(len(list(csv.reader(StringIO(text)))), 26)
